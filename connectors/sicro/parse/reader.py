from io import BytesIO

import openpyxl
import pandas as pd


def _normalize_cell(cell):
    # Header whitespace is inconsistent across files (e.g. "Preço Unitário
    # (R$)" vs "Preço Unitário\n(R$)"), so collapse all whitespace before
    # comparing.
    return " ".join(cell.split()) if isinstance(cell, str) else cell


def _normalize_row(row):
    return tuple(_normalize_cell(cell) for cell in row) if row else row


def _clean_string(value) -> str:
    # Basic standardization for string data columns (codigo/descricao/
    # unidade): collapse whitespace, strip, lowercase, so the same real-world
    # value never ends up as multiple distinct strings across files.
    text = "" if value is None else str(value)
    return " ".join(text.split()).strip().lower()


def _clean_numeric(value):
    # Some files use "-" as a placeholder for "no value" in numeric columns
    # instead of leaving the cell blank.
    if isinstance(value, str) and value.strip() == "-":
        return None
    return value

# Confirmed by opening real Relatório Sintético xlsx files from blob storage:
# single sheet, one row per item, but the header row (and its columns) differs
# per report type — Equipamentos has no unit price, only an hourly cost
# breakdown, and Mão de Obra has salário/encargos/periculosidade instead.
#
# Each report_type maps to a list of schema *variants* rather than a single
# schema: SICRO changed the Mão de Obra layout partway through 2025 (dropped
# Salário/Encargos Totais/Periculosidade, collapsed everything into a single
# "Custo (R$)" column), and older files under that report_type still use the
# original 7-column layout. New variants can be appended the same way if
# other report types change shape in the future.
REPORT_SCHEMAS = {
    "materiais": [
        {
            "header": ("Código", "Descrição", "Unidade", "Preço Unitário\n(R$)"),
            "columns": ("codigo", "descricao", "unidade", "preco_unitario"),
        },
    ],
    "equipamentos": [
        {
            "header": (
                "Código", "Descrição", "Valor de Aquisição (R$)", "Depreciação (R$/h)",
                "Oportunidade de Capital (R$/h)", "Seguros e Impostos (R$/h)",
                "Manutenção (R$/h)", "Operação (R$/h)", "Mão de Obra de Operação (R$/h)",
                "Custo Produtivo (R$/h)", "Custo Improdutivo (R$/h)",
            ),
            "columns": (
                "codigo", "descricao", "valor_aquisicao", "depreciacao",
                "oportunidade_capital", "seguros_impostos", "manutencao", "operacao",
                "mao_de_obra_operacao", "custo_produtivo", "custo_improdutivo",
            ),
        },
    ],
    "mao_de_obra": [
        {
            # Pre-2025 layout.
            "header": (
                "Código", "Descrição", "Unidade", "Salário (R$)", "Encargos Totais",
                "Custo (R$)", "Periculosidade/\nInsalubridade",
            ),
            "columns": (
                "codigo", "descricao", "unidade", "salario", "encargos_totais",
                "custo", "periculosidade_insalubridade",
            ),
        },
        {
            # 2025+ layout: salário/encargos/periculosidade columns were
            # dropped in favor of a single all-in cost column.
            "header": ("Código", "Descrição", "Unidade", "Custo\n (R$)"),
            "columns": ("codigo", "descricao", "unidade", "custo"),
        },
    ],
}


def parse_relatorio_sintetico(xlsx_bytes: bytes, report_type: str) -> pd.DataFrame:
    """Parses a Relatório Sintético xlsx into a flat DataFrame, using the
    header/column schema for the given report_type (materiais/equipamentos/
    mao_de_obra — see REPORT_SCHEMAS). Tries each known schema variant for
    the report_type in order, since the same report_type can have more than
    one layout depending on when the file was published. Raises ValueError
    if none match instead of silently skipping/guessing — a bad file becomes
    a 'failed' tracking row upstream, matching SicroClient._process_row's
    fail-loudly convention."""

    variants = [
        {**variant, "header": _normalize_row(variant["header"])}
        for variant in REPORT_SCHEMAS[report_type]
    ]

    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    rows = worksheet.iter_rows(values_only=True)

    # Files have a title/metadata banner (e.g. "CGCIT... SICRO...", region/
    # month, "Com desoneração") before the actual column header, so scan for
    # it instead of assuming it's row 1. Rows are also padded with trailing
    # None cells out to the worksheet's used width, so compare only the
    # leading cells that line up with each candidate header.
    matched = None
    for row in rows:
        normalized = _normalize_row(row)
        for variant in variants:
            if normalized[: len(variant["header"])] == variant["header"]:
                matched = variant
                break
        if matched:
            break
    else:
        expected = [variant["header"] for variant in variants]
        raise ValueError(
            f"None of the expected headers {expected!r} found in Relatório Sintético xlsx"
        )

    columns = matched["columns"]
    data = [row[: len(columns)] for row in rows if row and row[0] is not None]

    df = pd.DataFrame(data, columns=columns)

    string_columns = {"codigo", "descricao", "unidade"}
    for col in columns:
        if col in string_columns:
            df[col] = df[col].map(_clean_string)
        else:
            df[col] = df[col].map(_clean_numeric).astype(float)

    return df
